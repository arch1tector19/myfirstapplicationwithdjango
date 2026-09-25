from pathlib import Path

from django.core.files import File
from django.db import transaction
from openpyxl import load_workbook

from sbom.models import Component, UploadedFile


EXPECTED_HEADERS = [
    "Component",
    "Version",
    "Type",
    "BOM Reference",
    "PURL",
    "externalReferences",
    "Lang",
    "attack_surface",
    "security_function",
]


@transaction.atomic
def import_excel_file(file_path: str | Path) -> UploadedFile:

    Component.objects.all().delete()
   
    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    rows = worksheet.iter_rows(values_only=True)
    headers = list(next(rows, []))

    if headers != EXPECTED_HEADERS:
        workbook.close()
        raise ValueError(
            "Структура Excel-файла не соответствует ожидаемой. "
            f"Ожидались: {EXPECTED_HEADERS}"
        )

    components = []

    for row in rows:
        if not any(value is not None for value in row):
            continue

        values = list(row)

        while len(values) < len(EXPECTED_HEADERS):
            values.append(None)

        components.append(
            {
                "component": str(values[0] or ""),
                "version": str(values[1] or ""),
                "type": str(values[2] or ""),
                "bom_reference": str(values[3] or ""),
                "purl": str(values[4] or ""),
                "external_references": str(values[5] or ""),
                "lang": str(values[6] or ""),
                "attack_surface": str(values[7] or ""),
                "security_function": str(values[8] or ""),
            }
        )

    workbook.close()

    with open(file_path, "rb") as file:
        uploaded_file = UploadedFile.objects.create(
            original_name=file_path.name,
            file=File(file, name=file_path.name),
        )

    Component.objects.bulk_create(
        [
            Component(
                uploaded_file=uploaded_file,
                **component_data,
            )
            for component_data in components
        ]
    )

    return uploaded_file
